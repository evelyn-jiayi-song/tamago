#!/usr/bin/env python3
"""Build the companion speaker-script deck for Wobble Mechanisms."""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "simulator" / "docs"
DEFAULT_OUT = DOCS / "WOBBLE_MECHANISMS_SPEECH_SCRIPT.pptx"

BG = RGBColor(255, 255, 252)
PANEL = RGBColor(242, 244, 241)
TEXT = RGBColor(32, 37, 44)
MUTED = RGBColor(94, 103, 112)
RED = RGBColor(140, 21, 45)
PINK = RGBColor(204, 64, 112)
PURPLE = RGBColor(118, 67, 151)
BLUE = RGBColor(49, 91, 160)
CYAN = RGBColor(0, 126, 167)
GREEN = RGBColor(0, 132, 96)
AMBER = RGBColor(229, 132, 0)


SECTIONS = [
    ("WOBBLE MECHANISMS", "This project asks a small physical question: how can a bottom-heavy body turn a motor command into a readable, social gesture? I will trace one line from restoring torque and energy, through the electronics and firmware, into the dashboard and the final art direction.", RED),
    ("The motion loop", "The forward path is command, coil drive, shaft torque, reaction torque, and visible tilt. The dashboard observes orientation, inertial state, and motor state; stable-rock safety and peer filtering use those measurements, while ordinary run and rock commands remain open-loop.", CYAN),
    ("Electronics integration", "This is a repository wiring reference, not a manufactured PCB. The firmware test map is STEP 25, DIR 26, and ENN 27, while the schematic still needs reconciliation and labels a different pin relationship and a 10 volt supply. I keep that discrepancy visible, along with the capacitor, pull-up, motor reference, and unrecorded driver-current limit.", CYAN),
    ("Components inside one egg", "Each component has one job in the loop. The BNO055 senses, the ESP32 parses and schedules, the TMC2209 converts pulse and direction signals into configured phase current, the stepper produces shaft torque, and the mechanism turns reaction torque into body tilt. The sensor observes the body; generated steps are not yet an encoder measurement.", BLUE),
    ("Mechanical design: load path + center of mass", "The repository gives me a concrete starting model: 500 grams total mass, approximately 35 by 35 by 40 millimeter semi-axes, a 25 millimeter modeled CoM offset, a 35 millimeter contact radius, 0.1 millimeter surface roughness, and 2 millimeter shell thickness. The load path is motor and shaft to bearing support to offset mass and shell, then through the contact to the table. The bearing part and its measured radial or axial load are not yet validated, so I name them as the next mechanical check.", BLUE),
    ("Why the egg self-rights", "The restoring torque is negative because gravity pushes the body back toward upright. The model uses the contact-to-center-of-mass lever arm ell-CoM, while the 25 millimeter d-CoM value is a separate geometric model parameter. The model proxy m-g-d-CoM is 0.123 newton-meters per radian for the 500 gram, 25 millimeter example. IMU degrees must be converted to radians before using the equation.", AMBER),
    ("The rotational physics", "A force applied at a radius produces torque. Integrating torque over angle gives work; angular speed turns torque into mechanical power; a gear reduction trades speed for torque. The bottom line is the causal chain I will use when I talk about the actuator.", RED),
    ("Where the energy goes", "The energy story is electrical input, mechanical output, stored kinetic and gravitational energy, and then dissipation through friction, damping, coil heat, and sound. These equations define what the quantities mean. Electrical power, output torque, and loaded efficiency are still measurements for the next experiment, not claims of this prototype.", GREEN),
    ("Individual board communication", "First I establish one complete board contract: BNO055 orientation over I2C to the ESP32, then STEP, DIR, and ENN signals to the TMC2209, then phase current into the stepper. The current peer relationship is Egg A to the laptop, through filtering and delay, to Egg B. Both boards remain individually addressable; direct board-to-board wireless is the next step.", PURPLE),
    ("Peer coordination", "Peer detection turns gyro and dynamic acceleration into a filtered motion intent. The source score is amplified, clamped, and scaled by target fraction before mapping into speed, amplitude, and acceleration. Stale and stop detection exists in the prototype, but integrated delivery of that stop to the motor is not yet validated.", PURPLE),
    ("Codespace setup and rationale", "The repository is split into physics, firmware, dashboard, peer logic, and evidence rather than one long program. The shared interface is newline-delimited JSON for commands and telemetry. That lets me test the model without hardware, isolate transport from motion, and trace a run back to one command and one log.", AMBER),
    ("Runtime architecture", "At runtime, the laptop dashboard routes commands by board ID over two independent USB serial channels. Each board reports telemetry at about thirty samples per second over a 115200 bit-per-second JSON link. The peer path is laptop-mediated, so both boards remain independently addressable.", GREEN),
    ("Wireless next steps", "USB is the working path today. Wi-Fi is an experimental HTTP state, command, and heartbeat API, not yet a drop-in replacement: full USB command parity and automatic heartbeat servicing remain unfinished. The 2.4 gigahertz path needs its own power, hotspot, and heartbeat tests before it can carry the installation reliably.", PURPLE),
    ("Command surface", "These are the controls I can author: distance and amplitude in shaft revolutions, speed in revolutions per minute, acceleration in revolutions per minute per second, direction, cycles, tilt limit, intensity, and follower scale. Putting the units beside every field makes the interface a technical contract rather than a collection of unexplained sliders.", CYAN),
    ("Dashboard measurements", "The dashboard separates body measurements from generated motor state. Roll and pitch are in degrees; acceleration is meters per second squared; gyro is degrees per second; temperature is degrees Celsius. Position and target steps are software-generated state, not verified physical shaft motion.", AMBER),
    ("Physics vocabulary in the dashboard", "The vocabulary matters because raw sensor values are not yet meaning. Gyro magnitude is angular velocity, the accelerometer reports specific force including gravity, dynamic acceleration estimates motion beyond gravity, and taring makes yaw, roll, and pitch relative to a pose. Those features feed the peer detector.", PURPLE),
    ("Motion primitives and conversion", "The firmware has one-way travel, smooth travel with braking, rocking, and a tilt-limited over-tilt safety mode. Two hundred full steps per revolution times sixteen microsteps gives 3200 commanded steps per revolution. The 0.05 revolution example becomes 160 generated steps and an 18 degree commanded shaft angle; it is not encoder-verified.", GREEN),
    ("Smooth distance control", "A normal run can cut off at cruise speed. Smooth run calculates a braking speed from remaining steps, acceleration, and the steps-per-revolution constant, then limits the target speed before the endpoint. The 120 factor is an explicit unit-conversion factor, and the control law is designed to reduce endpoint impulse.", GREEN),
    ("Firmware realization", "The firmware path is parser, stepper state, PWM output, and telemetry. JSON fields become bounded pulse timing, direction, and enable signals; the same loop emits IMU and motor state. The configured ceilings are 1 to 240 revolutions per minute, 1 to 120 revolutions per minute per second, and 15,000 steps per second.", BLUE),
    ("Measured command-path evidence", "The current evidence is intentionally narrow: an unloaded USB command path. A 0.05 shaft-revolution command reached a software auto-stop at 160 generated steps; the log includes a 3.06 degree pitch peak, a 13.73 degree-per-second gyro peak, and a 110 to 113 millisecond USB command-to-idle response.", CYAN),
    ("What the current evidence measures", "The boundary is important. I can observe generated steps, IMU response, USB timing, and unloaded sensor temperature. I still need a marker or encoder for shaft angle, current and VMOT instrumentation, loaded tilt and natural frequency, repeatability, missed steps, and a judgment about whether peer delay feels intentional.", RED),
    ("Next experiment", "The next run closes the command-to-motion gap. I will measure the power rails, mark or encode shaft angle, add a known load and center-of-mass offset, and correlate commanded steps with shaft angle, body response, current, temperature, and missed steps. This turns a working command path into a measured mechanism.", AMBER),
    ("Making the wobble feel cute", "Cute is not a hidden motor property; it is an authoring hypothesis about timing and recovery. Small amplitude can feel contained, a slow ramp can feel soft, a short cycle can feel expressive, and delay can create social timing. The body trajectory is a measured response, so the formula and the mappings are starting points for motion studies, not validated perception results.", PINK),
    ("What I learned", "Most of my background is in software, digital and user-experience design, and art installations. This was the first time I soldered since my first year of college. I found the work incredibly rewarding and meditative, and I feel very blessed to have had the best mentors. Thank you to the lab for providing tools so generously; I definitely consumed a lot of 3D-printing resources.", CYAN),
    ("Photos and meals", "I left this page open for the evidence that does not fit in a schematic: the first soldering session, wiring and debugging, print iterations and resources, and meals or mentorship. Those images can make the build process and the lab context as legible as the final mechanism.", GREEN),
    ("Reproducibility", "The deck is generated from the repository source, the physical log, and rasterized diagram assets. The companion video is an explanatory visualization of one source egg and sixteen receivers with delayed, damped motion; it is not calibrated rigid-body dynamics and not evidence from sixteen hardware eggs. The deliverables keep the technical flow, the limitations, and the next experiment together.", RED),
]


def add_text(slide, value, x, y, w, h, size, color=TEXT, bold=False, align=PP_ALIGN.LEFT, font="Aptos"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, line in enumerate(str(value).split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(4)
        run = paragraph.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def build(output: Path = DEFAULT_OUT):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for index, (title, speech, accent) in enumerate(SECTIONS, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = BG
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.0), Inches(0.0), Inches(0.22), Inches(7.5))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.color.rgb = accent
        add_text(slide, f"{index:02d}", 0.72, 0.55, 0.55, 0.3, 14, accent, True)
        add_text(slide, title, 1.38, 0.45, 11.1, 0.58, 27, TEXT, True)
        add_text(slide, "COMPANION SPEECH SCRIPT", 1.4, 1.2, 4.0, 0.25, 13, MUTED, True)
        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.38), Inches(1.55), Inches(10.95), Inches(0.06))
        rule.fill.solid()
        rule.fill.fore_color.rgb = accent
        rule.line.color.rgb = accent
        add_text(slide, speech, 1.4, 2.0, 10.7, 3.7, 21, TEXT)
        add_text(slide, "Say it as a connected story; pause on the units and the validation boundary.", 1.4, 6.25, 10.7, 0.3, 13, accent, True)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build the Wobble Mechanisms speech-script deck")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    build(args.output)
